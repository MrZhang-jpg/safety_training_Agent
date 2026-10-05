# -*- coding: utf-8 -*-
"""
Word 导出模块：把结构化数据渲染为 .docx（python-docx）。

支持：
- exam_to_docx     : 考核试卷（考生卷 + 参考答案与解析页，支持下载留白作答）
- material_to_docx : 标准化培训资料（章节化，可直接打印下发）
- ledger_to_docx   : 安全培训台账（汇总 + 成绩明细表）

统一使用中文字体（标题黑体、正文宋体），保证 Word 打开不乱码。
"""
from __future__ import annotations

import io
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .schemas import Exam, MaterialSection, TrainingLedger, TrainingMaterial

# 中文字体
FONT_TITLE = "黑体"
FONT_BODY = "宋体"


def _set_run(run, cn_font: str = FONT_BODY, size: float = 10.5,
             bold: bool = False, color: Optional[str] = None) -> None:
    """设置 run 字体（含中文 eastAsia），避免 Word 中文乱码。"""
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _setup_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)


def _add_title(doc: Document, text: str, size: float = 16) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(text), cn_font=FONT_TITLE, size=size, bold=True)


def _add_para(doc: Document, text: str, size: float = 10.5,
              bold: bool = False, indent: bool = True,
              cn_font: str = FONT_BODY, align=None) -> None:
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    _set_run(p.add_run(text), cn_font=cn_font, size=size, bold=bold)
    return p


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    size = {0: 15, 1: 13, 2: 12}.get(level, 12)
    _set_run(p.add_run(text), cn_font=FONT_TITLE, size=size, bold=True)


def _add_blank_line(doc: Document, n: int = 1) -> None:
    for _ in range(n):
        doc.add_paragraph()


# =====================================================================
# 试卷
# =====================================================================
def exam_to_docx(exam: Exam, trainee_name: str = "") -> bytes:
    """渲染试卷 Word：考生卷 + 参考答案与解析（分页）。"""
    doc = Document()
    _setup_default_style(doc)

    _add_title(doc, exam.title, size=16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(sub.add_run(
        f"工种：{exam.craft_type}　教育层级：{exam.edu_level}　"
        f"满分：{exam.total_score:g}分　考试时间：90分钟"),
        size=10.5)
    _add_blank_line(doc)

    # 考生信息行
    info = doc.add_paragraph()
    _set_run(info.add_run(
        f"姓名：{trainee_name or '________'}　　单位/班组：________　　"
        f"日期：________　　得分：________"),
        size=10.5)
    _add_blank_line(doc)

    # -------- 试题部分 --------
    for q in exam.questions:
        head = doc.add_paragraph()
        _set_run(head.add_run(f"{q.index}.【{q.q_type}】（{q.score:g}分）"),
                 bold=True, size=11)
        body = doc.add_paragraph()
        body.paragraph_format.first_line_indent = Pt(21)
        _set_run(body.add_run(q.content), size=11)
        for opt in (q.options or []):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Pt(24)
            _set_run(op.add_run(opt), size=10.5)
        # 简答题留作答空白
        if q.q_type == "简答题":
            _add_blank_line(doc, 3)
        else:
            _add_blank_line(doc, 1)

    # 页末提示
    tip = doc.add_paragraph()
    _set_run(tip.add_run("—— 试卷作答完毕，请将试卷上交；以下为参考答案与解析，请勿提前翻阅 ——"),
             size=10, color="808080")
    doc.add_page_break()

    # -------- 参考答案与解析 --------
    _add_title(doc, "参考答案与解析", size=15)
    _add_blank_line(doc)
    for q in exam.questions:
        p = doc.add_paragraph()
        _set_run(p.add_run(f"{q.index}.【{q.q_type}】答案：{q.answer or '—'}"),
                 bold=True, size=10.5)
        if q.analysis:
            a = doc.add_paragraph()
            a.paragraph_format.left_indent = Pt(18)
            _set_run(a.add_run(f"解析：{q.analysis}"), size=10)

    # 参考来源
    if exam.references:
        _add_blank_line(doc)
        _add_heading(doc, "参考资料来源", level=2)
        for src in sorted(set(exam.references)):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            _set_run(p.add_run(f"· {src}"), size=9.5, color="606060")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =====================================================================
# 培训资料
# =====================================================================
def material_to_docx(material: TrainingMaterial) -> bytes:
    """渲染标准化培训资料 Word（章节化）。"""
    doc = Document()
    _setup_default_style(doc)

    _add_title(doc, material.title, size=16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(sub.add_run(
        f"适用对象：{material.craft_type}　教育层级：{material.edu_level}"),
        size=10.5)
    _add_blank_line(doc)

    if material.summary:
        _add_heading(doc, "一、培训目的与适用范围", level=1)
        for para in str(material.summary).split("\n"):
            if para.strip():
                _add_para(doc, para.strip())

    for i, sec in enumerate(material.sections, start=2):
        _add_heading(doc, f"{_cn_num(i)}、{sec.heading}", level=1)
        for para in str(sec.content).split("\n"):
            if para.strip():
                _add_para(doc, para.strip())

    if material.references:
        _add_blank_line(doc)
        _add_heading(doc, "参考资料来源", level=2)
        for src in sorted(set(material.references)):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            _set_run(p.add_run(f"· {src}"), size=9.5, color="606060")

    _add_blank_line(doc)
    tail = doc.add_paragraph()
    _set_run(tail.add_run("本资料由安全培训智能 Agent 依据企业内部规程/交底/题库自动生成，"
                          "仅供三级安全教育与岗前培训使用，需安全员复核确认。"),
             size=9.5, color="808080")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _cn_num(i: int) -> str:
    mapping = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五",
               6: "六", 7: "七", 8: "八", 9: "九", 10: "十"}
    return mapping.get(i, str(i))


# =====================================================================
# 台账
# =====================================================================
def ledger_to_docx(ledger: TrainingLedger) -> bytes:
    """渲染安全培训台账 Word（汇总信息 + 成绩明细表）。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    _setup_default_style(doc)

    _add_title(doc, "安全培训台账", size=16)
    _add_blank_line(doc)

    rows = [
        ("培训主题", ledger.training_topic or f"{ledger.craft_type}{ledger.edu_level}三级安全教育"),
        ("工种", ledger.craft_type),
        ("教育层级", ledger.edu_level),
        ("主讲人", ledger.trainer or "________"),
        ("培训地点", ledger.location or "________"),
        ("培训日期", ledger.training_date or "________"),
        ("培训时长", f"{ledger.duration_hours:g} 小时"),
        ("参训人数", str(ledger.trainee_count)),
        ("平均得分", f"{ledger.average_score:g} 分"),
        ("合格人数", str(ledger.pass_count)),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        _set_run(c0.paragraphs[0].add_run(k), bold=True, size=10.5)
        _set_run(c1.paragraphs[0].add_run(v), size=10.5)

    _add_blank_line(doc)
    _add_heading(doc, "参训人员成绩明细", level=1)

    names = list(ledger.trainee_names) or []
    scores = {g.trainee_name: g for g in ledger.grade_results if g.trainee_name}
    detail = doc.add_table(rows=1 + len(names), cols=4)
    detail.style = "Table Grid"
    header = detail.rows[0].cells
    for j, h in enumerate(["序号", "姓名", "得分", "是否合格"]):
        _set_run(header[j].paragraphs[0].add_run(h), bold=True, size=10.5)
    for i, name in enumerate(names, start=1):
        cells = detail.rows[i].cells
        g = scores.get(name)
        score = f"{g.total_score:g} 分" if g else "________"
        passed = "合格" if g and g.passed else ("不合格" if g else "________")
        _set_run(cells[0].paragraphs[0].add_run(str(i)), size=10.5)
        _set_run(cells[1].paragraphs[0].add_run(name), size=10.5)
        _set_run(cells[2].paragraphs[0].add_run(score), size=10.5)
        _set_run(cells[3].paragraphs[0].add_run(passed), size=10.5)

    _add_blank_line(doc)
    if ledger.grade_summary:
        _add_heading(doc, "培训小结", level=1)
        _add_para(doc, ledger.grade_summary)
    tail = doc.add_paragraph()
    _set_run(tail.add_run(ledger.remark or ""), size=9.5, color="808080")
    _add_blank_line(doc)
    sign = doc.add_paragraph()
    _set_run(sign.add_run("安全员签字：____________　　日期：____________"), size=10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
