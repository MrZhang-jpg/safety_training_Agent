# -*- coding: utf-8 -*-
"""
多格式 EHS 文档加载器（面向对象）。

支持：
- .docx / .doc : Word（.doc 在 Windows 下用 Word COM 转 .docx，Linux 下用 LibreOffice）
- .pdf          : pypdf 逐页提取
- .pptx / .ppt  : PowerPoint（.ppt 同样需要转换）
- .xlsx / .xls  : 表格按行结构化为文本
- .txt          : 直接读取

老格式 .doc/.ppt 转换结果会缓存到 data_cache/converted，避免重复转换。
输出统一为 LangChain Document，metadata 记录来源文件、所属知识类别、页码/工作表。
"""
from __future__ import annotations

import hashlib
import os
import platform
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Optional

from langchain_core.documents import Document

from .config import settings


class LegacyConverter:
    """把老式 .doc/.ppt 转换为现代 .docx/.pptx，带缓存。"""

    def __init__(self, cache_dir: Path):
        self.cache_dir = cache_dir
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        self.is_windows = platform.system().lower().startswith("win")

    def _cache_path(self, src: Path, target_ext: str) -> Path:
        key = hashlib.md5(str(src.resolve()).encode("utf-8")).hexdigest()[:12]
        return self.cache_dir / f"{src.stem}_{key}{target_ext}"

    def convert(self, src: Path, target_ext: str) -> Optional[Path]:
        dst = self._cache_path(src, target_ext)
        if dst.exists() and dst.stat().st_size > 0:
            return dst

        if self.is_windows:
            method = self._convert_win_word if target_ext == ".docx" else self._convert_win_ppt
            try:
                if method(src, dst):
                    return dst
            except Exception as exc:  # COM 失败则尝试 LibreOffice
                print(f"[WARN] COM 转换 {src.name} 失败: {exc}")

        if self._convert_soffice(src, dst):
            return dst
        return None

    # ---------- Windows COM ----------
    def _convert_win_word(self, src: Path, dst: Path) -> bool:
        import pythoncom
        import win32com.client as win32

        for attempt in (1, 2):
            word = None
            try:
                pythoncom.CoInitialize()
                word = win32.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                try:
                    word.AutomationSecurity = 3  # 禁用宏，避免转换时弹窗
                except Exception:
                    pass
                doc = word.Documents.Open(
                    FileName=str(src.resolve()), ConfirmConversions=False,
                    ReadOnly=True, AddToRecentFiles=False, PasswordDocument="",
                    Visible=False, NoEncodingDialog=True,
                )
                doc.SaveAs(FileName=str(dst.resolve()), FileFormat=16)  # .docx
                doc.Close(SaveChanges=False)
                return dst.exists()
            except Exception as exc:
                print(f"[WARN] Word COM 第 {attempt} 次转换失败 {src.name}: {exc}")
                if word is not None:
                    try:
                        word.Quit()
                    except Exception:
                        pass
            finally:
                pythoncom.CoUninitialize()
        return False

    def _convert_win_ppt(self, src: Path, dst: Path) -> bool:
        import pythoncom
        import win32com.client as win32

        app = None
        try:
            pythoncom.CoInitialize()
            app = win32.DispatchEx("PowerPoint.Application")
            pres = app.Presentations.Open(str(src.resolve()), WithWindow=False)
            pres.SaveAs(str(dst.resolve()), 24)  # 24 = ppSaveAsOpenXMLPresentation
            pres.Close()
            return dst.exists()
        except Exception:
            return False
        finally:
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()

    # ---------- LibreOffice（Linux / Docker / COM 失败兜底）----------
    def _convert_soffice(self, src: Path, dst: Path) -> bool:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            return False
        out_dir = self.cache_dir
        cmd = [soffice, "--headless", "--convert-to", dst.suffix.lstrip("."),
               "--outdir", str(out_dir), str(src)]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=180)
        except Exception as exc:
            print(f"[WARN] LibreOffice 转换 {src.name} 失败: {exc}")
            return False
        produced = out_dir / f"{src.stem}{dst.suffix}"
        if produced.exists() and produced != dst:
            shutil.move(str(produced), str(dst))
        return dst.exists()


class MultiFormatLoader:
    """遍历数据目录并把各类 EHS 文档加载为 Document。"""

    def __init__(self):
        self.data_dir = settings.data_path
        self.converter = LegacyConverter(settings.converted_path)
        self.stats = {"loaded_files": 0, "failed_files": 0, "skipped": 0}

    # ---------- 文件遍历 ----------
    def iter_files(self) -> list[Path]:
        if not self.data_dir.exists():
            raise FileNotFoundError(f"数据目录不存在: {self.data_dir}")
        files: list[Path] = []
        for path in sorted(self.data_dir.rglob("*")):
            if not path.is_file():
                continue
            ext = path.suffix.lower()
            if ext not in settings.supported_extensions:
                self.stats["skipped"] += 1
                continue
            if ext in settings.excluded_extensions:
                self.stats["skipped"] += 1
                continue
            if settings.include_dirs:
                rel_parts = path.relative_to(self.data_dir).parts
                if not any(part in settings.include_dirs for part in rel_parts):
                    self.stats["skipped"] += 1
                    continue
            files.append(path)
        return files

    # ---------- 元数据 ----------
    def _base_meta(self, path: Path) -> dict:
        rel = path.relative_to(self.data_dir)
        parts = rel.parts
        return {
            "source": rel.as_posix(),
            "file_name": path.name,
            "category": parts[0] if len(parts) > 1 else "根目录",
        }

    # ---------- 单文件分发 ----------
    def load_file(self, path: Path) -> list[Document]:
        ext = path.suffix.lower()
        if ext == ".docx":
            return self._load_docx(path)
        if ext == ".doc":
            meta = self._base_meta(path)
            converted = self.converter.convert(path, ".docx")
            return self._load_docx(converted, meta) if converted else []
        if ext == ".pdf":
            return self._load_pdf(path)
        if ext == ".pptx":
            return self._load_pptx(path)
        if ext == ".ppt":
            meta = self._base_meta(path)
            converted = self.converter.convert(path, ".pptx")
            return self._load_pptx(converted, meta) if converted else []
        if ext == ".xlsx":
            return self._load_xlsx(path)
        if ext == ".xls":
            return self._load_xls(path)
        if ext == ".txt":
            return self._load_txt(path)
        return []

    # ---------- 各格式具体实现 ----------
    def _load_docx(self, path: Path, meta_override: Optional[dict] = None) -> list[Document]:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        chunks: list[str] = []

        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                chunks.append(t)

        for ti, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                chunks.append(f"【表格{ti}】\n" + "\n".join(rows))

        if not chunks:
            return []
        meta = meta_override or self._base_meta(path)
        return [Document(page_content="\n".join(chunks), metadata=meta)]

    def _load_pdf(self, path: Path) -> list[Document]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        docs: list[Document] = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = (page.extract_text() or "").strip()
            except Exception:
                text = ""
            if text:
                meta = self._base_meta(path)
                meta["page"] = i
                docs.append(Document(page_content=text, metadata=meta))
        return docs

    def _load_pptx(self, path: Path, meta_override: Optional[dict] = None) -> list[Document]:
        from pptx import Presentation

        prs = Presentation(str(path))
        chunks: list[str] = []
        for si, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        texts.append(t)
                if shape.has_table:
                    for row in shape.table.rows:
                        line = " | ".join(c.text.strip() for c in row.cells)
                        if line.strip(" |"):
                            texts.append(line)
            if texts:
                chunks.append(f"【幻灯片 {si}】\n" + "\n".join(texts))
        if not chunks:
            return []
        meta = meta_override or self._base_meta(path)
        return [Document(page_content="\n".join(chunks), metadata=meta)]

    def _load_xlsx(self, path: Path) -> list[Document]:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        docs: list[Document] = []
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            if lines:
                meta = self._base_meta(path)
                meta["sheet"] = ws.title
                docs.append(Document(page_content=f"【工作表 {ws.title}】\n" + "\n".join(lines),
                                     metadata=meta))
        wb.close()
        return docs

    def _load_xls(self, path: Path) -> list[Document]:
        import xlrd

        book = xlrd.open_workbook(str(path))
        docs: list[Document] = []
        for ws in book.sheets():
            lines = []
            for r in range(ws.nrows):
                cells = [str(v).strip() for v in ws.row_values(r)]
                if any(cells):
                    lines.append(" | ".join(cells))
            if lines:
                meta = self._base_meta(path)
                meta["sheet"] = ws.name
                docs.append(Document(page_content=f"【工作表 {ws.name}】\n" + "\n".join(lines),
                                     metadata=meta))
        return docs

    def _load_txt(self, path: Path) -> list[Document]:
        for enc in ("utf-8", "gbk", "gb18030"):
            try:
                text = path.read_text(encoding=enc).strip()
                break
            except UnicodeDecodeError:
                continue
        else:
            return []
        if not text:
            return []
        return [Document(page_content=text, metadata=self._base_meta(path))]

    # ---------- 全量加载 ----------
    def load_all(self, show_progress: bool = True) -> list[Document]:
        files = self.iter_files()
        all_docs: list[Document] = []
        total = len(files)
        for i, path in enumerate(files, start=1):
            try:
                docs = self.load_file(path)
                if docs:
                    all_docs.extend(docs)
                    self.stats["loaded_files"] += 1
                else:
                    self.stats["failed_files"] += 1
            except Exception as exc:
                self.stats["failed_files"] += 1
                print(f"[WARN] 加载失败 {path.name}: {exc}")
            if show_progress and (i % 10 == 0 or i == total):
                print(f"  文档加载进度 {i}/{total}")
        print(f"文档加载完成: 成功 {self.stats['loaded_files']}，"
              f"失败 {self.stats['failed_files']}，跳过 {self.stats['skipped']}")
        return all_docs
